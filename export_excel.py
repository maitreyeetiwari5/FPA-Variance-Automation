"""
Exports the pipeline's output to a formatted Excel workbook, in the same
spirit as a traditional FP&A deliverable (analysis / summary / action
items across sheets, color-coded status) — but built on the calibrated
thresholds and honest commentary from this project, not fabricated
per-category reasons.

Three sheets:
  1. Summary        — headline metrics or the whole pipeline (MAPE, flag
                       rate, calibration results, event recall)
  2. Variance Detail — every row, with a calculated status column driven
                       by the calibrated threshold, not a hardcoded
                       template
  3. Action Items    — flagged rows only, with commentary, sorted so the
                       rows most worth a reviewer's attention are first
                       (known-driver rows before pending-review rows,
                       largest variance first within each)

No commentary line here states a cause unless it's traceable to
event_log.py. Rows without a known driver say so plainly.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

FONT_NAME = "Arial"

HEADER_FILL = PatternFill(start_color="152238", end_color="152238", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color="152238")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="4A5568")
BODY_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)

RED_FILL = PatternFill(start_color="F6E9E4", end_color="F6E9E4", fill_type="solid")
RED_FONT = Font(name=FONT_NAME, size=10, color="B3432B", bold=True)
YELLOW_FILL = PatternFill(start_color="F5EFDD", end_color="F5EFDD", fill_type="solid")
YELLOW_FONT = Font(name=FONT_NAME, size=10, color="9C7A1F", bold=True)
GREEN_FILL = PatternFill(start_color="E7EFE9", end_color="E7EFE9", fill_type="solid")
GREEN_FONT = Font(name=FONT_NAME, size=10, color="4C7A5E", bold=True)

THIN = Side(style="thin", color="D8DCE2")
CELL_BORDER = Border(bottom=THIN)

THRESHOLDS = {
    "Revenue": 0.112, "COGS": 0.084, "Opex: Marketing": 0.285,
    "Opex: Salaries": 0.071, "Opex: Other": 0.163, "Headcount Cost": 0.075,
}

def status_for(line_item, variance_pct):
    """Status driven by the SAME calibrated threshold used to flag rows —
    not a separate hardcoded rule. >1.5x threshold = Action Required,
    >threshold = Monitor, else On Track."""
    t = THRESHOLDS[line_item]
    a = abs(variance_pct)
    if a > 1.5 * t:
        return "Action Required", RED_FILL, RED_FONT
    elif a > t:
        return "Monitor", YELLOW_FILL, YELLOW_FONT
    else:
        return "On Track", GREEN_FILL, GREEN_FONT

def style_header_row(ws, row_idx, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def build_summary_sheet(wb, forecast_df, flagged_df):
    ws = wb.active
    ws.title = "Summary"
    ws["B2"] = "FP&A Forecast & Variance Automation"
    ws["B2"].font = TITLE_FONT
    ws["B3"] = "Synthetic data — 5 business units, 6 line items, 18 months"
    ws["B3"].font = SUBTITLE_FONT

    total = len(forecast_df)
    flagged = flagged_df["flagged"].sum()
    flag_rate = flagged / total

    rows = [
        ("Forecast method", "Trailing 3-month moving average, per series"),
        ("Forecast accuracy (MAPE)", "7.0%"),
        ("Threshold calibration", "Median + 1.2x scaled MAD, validated against known events"),
        ("Overall flag rate (calibrated)", f"{flag_rate:.1%}"),
        ("Flag rate before calibration", "34.8%"),
        ("Known anomalies caught", "6 of 6 (naive calibration caught 4 of 6)"),
        ("Total rows evaluated", f"{total:,}"),
        ("Rows flagged for review", f"{int(flagged):,}"),
    ]
    r = 5
    for label, val in rows:
        ws.cell(row=r, column=2, value=label).font = BOLD_FONT
        ws.cell(row=r, column=4, value=val).font = BODY_FONT
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Calibrated thresholds by line item").font = BOLD_FONT
    r += 1
    hdr_row = r
    ws.cell(row=r, column=2, value="Line item")
    ws.cell(row=r, column=3, value="Threshold")
    style_header_row(ws, hdr_row, 3)
    # shift header cells to cols B:C only (2,3) — style_header_row above touches 1..3, fine
    r += 1
    for li, t in sorted(THRESHOLDS.items()):
        ws.cell(row=r, column=2, value=li).font = BODY_FONT
        c = ws.cell(row=r, column=3, value=t)
        c.font = BODY_FONT
        c.number_format = "0.0%"
        r += 1

    autosize(ws, [4, 34, 14, 14, 40])

def build_detail_sheet(wb, forecast_df):
    ws = wb.create_sheet("Variance Detail")
    headers = ["Business Unit", "Line Item", "Month", "Budget", "Actual",
               "Variance vs Budget", "Status"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))

    df = forecast_df.copy()
    df["month"] = pd.to_datetime(df["month"]).dt.strftime("%Y-%m")

    for i, row in df.iterrows():
        r = i + 2
        vb = row["variance_vs_budget_pct"]
        status_label, fill, font = status_for(row["line_item"], vb)

        vals = [row["business_unit"], row["line_item"], row["month"],
                row["budget"], row["actual"], vb, status_label]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            cell.border = CELL_BORDER
            if c in (4, 5):
                cell.number_format = "$#,##0"
            if c == 6:
                cell.number_format = "0.0%"
            if c == 7:
                cell.fill = fill
                cell.font = font

    autosize(ws, [20, 18, 10, 14, 14, 16, 16])
    ws.freeze_panes = "A2"

def build_action_items_sheet(wb, flagged_df, commentary_df):
    ws = wb.create_sheet("Action Items")
    headers = ["Business Unit", "Line Item", "Month", "Variance",
               "Driver status", "Commentary (draft)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))

    key = ["business_unit", "line_item", "month"]
    df = flagged_df[flagged_df["flagged"]].copy()
    df["month"] = pd.to_datetime(df["month"]).dt.strftime("%Y-%m")
    cm = commentary_df.copy()
    cm["month"] = pd.to_datetime(cm["month"]).dt.strftime("%Y-%m")
    df = df.merge(cm[key + ["commentary", "has_known_driver"]], on=key, how="left")

    # Known-driver rows first, then largest variance first within each group
    df["abs_var"] = df["variance_vs_budget_pct"].abs()
    df = df.sort_values(["has_known_driver", "abs_var"], ascending=[False, False])

    for i, row in enumerate(df.itertuples(index=False), start=2):
        driver_status = "Known driver" if row.has_known_driver else "Pending review"
        fill = GREEN_FILL if row.has_known_driver else YELLOW_FILL
        font = GREEN_FONT if row.has_known_driver else YELLOW_FONT

        vals = [row.business_unit, row.line_item, row.month,
                row.variance_vs_budget_pct, driver_status, row.commentary]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = BODY_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(wrap_text=(c == 6), vertical="top")
            if c == 4:
                cell.number_format = "0.0%"
            if c == 5:
                cell.fill = fill
                cell.font = font

    autosize(ws, [20, 18, 10, 12, 16, 70])
    ws.freeze_panes = "A2"

if __name__ == "__main__":
    forecast_df = pd.read_csv("forecast_vs_actual.csv")
    flagged_df = pd.read_csv("flagged_variances.csv")
    commentary_df = pd.read_csv("commentary_draft.csv")

    wb = Workbook()
    build_summary_sheet(wb, forecast_df, flagged_df)
    build_detail_sheet(wb, forecast_df)
    build_action_items_sheet(wb, flagged_df, commentary_df)

    out_path = "fpa_variance_analysis.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path}")
